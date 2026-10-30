"""The Capacity Fee Workbook's structure and data, read back with openpyxl
(no Excel needed; the figures are checked in Excel by test_workbook_excel.py)."""

from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

import build_workbook as cli
from core import fee, money, periods, points, tariffs, workbook
from core.tariffs import TSOS
from core.workbook import WorkbookError, build_workbook, stretches

from .helpers import make_frame, make_row, make_table

BUILD_DAY = date(2026, 10, 29)


@pytest.fixture(scope="module")
def real():
    table = tariffs.load_tariffs(tariffs.DEFAULT_PATH).table
    catalog = points.build_catalog(table)
    return table, catalog, build_workbook(table, catalog, "tariffs_real.xlsx", BUILD_DAY)


@pytest.fixture(scope="module")
def saved(real, tmp_path_factory):
    path = tmp_path_factory.mktemp("wb") / "workbook.xlsx"
    real[2].save(path)
    return load_workbook(path)


def column(ws, col: int) -> list:
    return [ws.cell(row, col).value for row in range(2, ws.max_row + 1)]


def named_cell(wb, name: str):
    (sheet, ref), = wb.defined_names[name].destinations
    return wb[sheet][ref.replace("$", "")]


# --- 2. data sheets ---------------------------------------------------------------------------------

def test_rules_follow_the_app(saved):
    rules = {row[0]: row for row in saved["Rules"].iter_rows(min_row=2, values_only=True)}
    assert set(rules) == set(TSOS)
    assert rules["GCA"][5] == 4 and all(rules[c][5] is None for c in TSOS if c != "GCA")
    assert rules["Gastran"][6:8] == ("no", "yes")
    for code in ("FGSZ", "EUSTR", "GCA", "NET4GAS"):
        assert rules[code][6:8] == ("yes", "yes")
    assert rules["FGSZ"][2:5] == ("HUF", "kWh/h", 0) and rules["BGTRGAZ"][3] == "kWh/d"


def test_stretches_follow_the_superseding_rule():
    early = make_row(valid_from="2025-01-01", valid_to=None, prices={"Year": 1.0})
    middle = make_row(valid_from="2026-01-01", valid_to="2026-12-31", prices={"Year": 2.0})
    table = make_table(early, middle)
    key = ("FGSZ", "21Z000000000003C", "Entry")
    got = [(first, last, row.year_price()) for first, last, row in stretches(table, key)]
    assert got == [(date(2025, 1, 1), date(2025, 12, 31), 1), (date(2026, 1, 1), date(2026, 12, 31), 2),
                   (date(2027, 1, 1), workbook.OPEN_END, 1)]


def test_stretches_leave_out_gaps():
    table = make_table(make_row(valid_from="2025-10-01", valid_to="2025-12-31"),
                       make_row(valid_from="2026-02-01", valid_to="2026-03-31"))
    got = [(first, last) for first, last, _ in stretches(table, ("FGSZ", "21Z000000000003C", "Entry"))]
    assert got == [(date(2025, 10, 1), date(2025, 12, 31)), (date(2026, 2, 1), date(2026, 3, 31))]


def test_stretches_match_product_parts_on_the_real_file(real):
    table, catalog, _ = real
    for point in (*catalog.singles, *catalog.lng):
        parts = stretches(table, point.key)
        for first, last, row in parts:
            end = min(last, first + timedelta(days=400))
            assert [p.row for p in table.product_parts(*point.key, first, end)] == [row]


def test_points_and_tariffs_sheets(real, saved):
    table, catalog, _ = real
    points_ws, tariffs_ws = saved["Points"], saved["Tariffs"]
    all_points = [*catalog.singles, *catalog.lng]
    assert len(column(points_ws, 1)) == len(all_points) == 123
    moson = next(p for p in catalog.singles if p.tso == "GCA" and p.fee_label().startswith("Mosonmagyaróvár AT>HU"))
    row = 2 + all_points.index(moson)
    first, count = points_ws.cell(row, 10).value, points_ws.cell(row, 11).value
    expected = stretches(table, moson.key)
    assert count == len(expected)
    for i, (start, end, tariff_row) in enumerate(expected):
        cells = [c.value for c in tariffs_ws[first + i]]
        assert cells[0] == row - 1
        assert (cells[1].date(), cells[2].date(), cells[3].date()) == (start, end, tariff_row.valid_from)
        assert cells[4] == (1 if tariff_row.open_ended else 0)
        assert cells[6:] == [float(tariff_row.prices[c]) for c in periods.PRICE_COLUMNS]


def test_too_many_stretches_stop_the_build(monkeypatch, real):
    table, catalog, _ = real
    monkeypatch.setattr(workbook, "MAX_STRETCH_SLOTS", 1)
    with pytest.raises(WorkbookError, match="room for 1"):
        build_workbook(table, catalog, "x", BUILD_DAY)


def test_lists_sheet(real, saved):
    _, catalog, _ = real
    lists = saved["Lists"]
    labels = column(lists, 10)
    assert labels[:43] == [b.label() for b in catalog.bundles] and len(catalog.bundles) == 43
    assert labels[43:57] == [p.lng_label() for p in catalog.lng]
    assert len(labels) == 43 + 14 + 109
    tsos = {name: (start, count) for name, start, count in
            zip(column(lists, 5), column(lists, 7), column(lists, 8)) if name}
    start, count = tsos["Gastrans"]
    assert count == 4
    fgsz_start, fgsz_count = tsos["FGSZ"]
    fgsz = labels[fgsz_start - 2:fgsz_start - 2 + fgsz_count]
    assert "Kiskundorozsma 2 HU>RS (21Z000000000505P) - Interruptible" in fgsz
    assert column(lists, 15)[:4] == ["CZK", "HUF", "PLN", "RON"]


# --- 4. the Calculator sheet --------------------------------------------------------------------------

def test_only_the_calculator_is_visible_and_protected(saved):
    assert [ws.title for ws in saved.worksheets if ws.sheet_state == "visible"] == ["Calculator"]
    calculator = saved["Calculator"]
    assert calculator.protection.sheet and not calculator.protection.password
    for name in ("in_booking", "in_tso", "in_point", "in_instrument", "in_period", "in_day", "in_capacity",
                 "in_unit", "in_rate_HUF"):
        assert named_cell(saved, name).protection.locked is False, name
    assert named_cell(saved, "out1_total").protection.locked is not False


def test_defaults(saved):
    values = {name: named_cell(saved, name).value for name in (
        "in_booking", "in_tso", "in_point", "in_instrument", "in_period", "in_capacity", "in_unit", "in_rate_HUF")}
    assert values == {"in_booking": "Bundled", "in_tso": "FGSZ",
                      "in_point": "Mosonmagyaróvár AT>HU (Gas Connect Austria + FGSZ)", "in_instrument": "Quarter",
                      "in_period": "2026 Q4", "in_capacity": None, "in_unit": "kWh/h", "in_rate_HUF": None}
    assert named_cell(saved, "in_day").value.date() == date(2026, 12, 5)
    assert named_cell(saved, "in_period").number_format == "@"


def test_dropdowns(saved):
    calculator = saved["Calculator"]
    sources = {}
    for dv in calculator.data_validations.dataValidation:
        for ref in str(dv.sqref).split():
            sources[ref] = (dv.type, dv.formula1)
    cell = lambda name: named_cell(saved, name).coordinate
    assert sources[cell("in_booking")] == ("list", "lst_booking")
    assert sources[cell("in_tso")] == ("list", "lst_tso_name")
    assert sources[cell("in_point")] == ("list", "dv_points")
    assert sources[cell("in_instrument")] == ("list", "lst_instrument")
    assert sources[cell("in_period")] == ("list", "dv_periods")
    assert sources[cell("in_unit")] == ("list", "lst_unit")
    assert sources[cell("in_day")][0] == "date"
    assert saved.defined_names["dv_points"].attr_text.startswith("OFFSET(")


def test_every_named_output_exists(saved):
    names = ["msg_main", "out_heading", "out_eur_total", "out_eur_caption", "out_eur_msg", "in_rates"] + [
        f"out{n}_{field}" for n in (1, 2) for field in (
            "heading", "tso", "direction", "capacity", "price", "total", "tariff_used", "caption", "rounding",
            "interruptible", "equal", "months", "days", "amounts", "invoice_total")]
    for name in names:
        assert name in saved.defined_names, name
    assert not any(name.startswith(("s1_", "s2_")) for name in saved.defined_names)  # would read as cell S1


def test_formulas_are_language_independent_and_self_contained(saved, real):
    for ws in saved.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    assert "TEXT(" not in cell.value.upper().replace("TEXTJOIN(", ""), cell.coordinate
                    assert "[" not in cell.value, cell.coordinate  # no links to other files
    assert saved.vba_archive is None and not saved._external_links
    assert real[2]["Calculator"]["A2"].value.startswith("Built from tariffs_real.xlsx on 2026-10-29")


# --- 5. the command line ------------------------------------------------------------------------------

def test_build_from_the_sample_file(tmp_path, capsys):
    output = tmp_path / "out" / "fee.xlsx"
    assert cli.main([str(tariffs.DEFAULT_SAMPLE_PATH), "-o", str(output)]) == 0
    wb = load_workbook(output)
    assert wb["Calculator"]["A2"].value.startswith("Built from tariffs_sample.xlsx on ")
    assert "Wrote" in capsys.readouterr().out


def test_a_rejected_file_writes_nothing(tmp_path, capsys):
    frame = make_frame(make_row()).drop(columns=["Year"])
    source = tmp_path / "no_year.xlsx"
    frame.to_excel(source, sheet_name=tariffs.SHEET_NAME, index=False)
    output = tmp_path / "fee.xlsx"
    assert cli.main([str(source), "-o", str(output)]) == 1
    assert not output.exists()
    expected = tariffs.load_tariffs(source).errors
    assert expected and expected[0] in capsys.readouterr().err
