"""Capacity Fee Workbook: tab 2 as one Excel file, built from the tariff data.

`build_workbook` returns an openpyxl workbook with formulas only (no macros).
Everything the app resolves in Python is resolved here and written as data:
the point catalog and its labels, the stretches of days each tariff row
applies to, and each TSO's rules. Excel then only looks up, overlaps dates,
rounds and splits. Every input and result cell has a defined name (`in_*`,
`out_*`, `msg_main`), which is how the tests address them.

Formulas avoid spilling functions (openpyxl cannot write them) and `TEXT()`
(its format codes depend on the Excel language)."""

from __future__ import annotations

from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.formatting.rule import Rule
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from . import fee, money, periods
from .points import EXIT, Catalog, Point, tso_name
from .tariffs import HOURS_PER_UNIT, INTERRUPTIBLE, TSOS, TariffRow, TariffTable

BOOKING_TYPES = ("Bundled", "Single side", "LNG entry")
DEFAULT_BUNDLE = "Mosonmagyaróvár AT>HU (Gas Connect Austria + FGSZ)"
DEFAULT_TSO = "FGSZ"
DEFAULT_PERIOD = "2026 Q4"
DEFAULT_DAY = date(2026, 12, 5)
OPEN_END = date(9999, 12, 31)  # the last day Excel knows: an open-ended stretch runs to it
MIN_STRETCH_SLOTS = 8
MAX_STRETCH_SLOTS = 40
SIDE_HEIGHT = 26  # rows of one side's results on the Calculator sheet
FIRST_SIDE_ROW = 15
# Side 1 (a single point, or a bundle's exit) and side 2 (a bundle's entry). Names like "s1_x" would
# read as cell S1 in Excel, so the sides are "sa" and "sb".
SIDE_PREFIX = {1: "sa_", 2: "sb_"}

INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
GREY = DifferentialStyle(fill=PatternFill("solid", bgColor="E7E6E6"), font=Font(color="A6A6A6"))
HEADER_FONT = Font(bold=True)
NOTE_FONT = Font(italic=True, color="595959")


class WorkbookError(ValueError):
    """The tariff data does not fit the workbook layout; the message says why."""


# --- data prepared in Python ------------------------------------------------------------------

def stretches(table: TariffTable, key: tuple[str, str, str]) -> list[tuple[date, date, TariffRow]]:
    """The point's days split where the applicable row changes, as (first, last, row),
    in date order and without overlaps; days no row applies to are left out. The
    row on each day is `TariffTable.find_row`'s, so a later row supersedes an
    earlier one exactly as in `product_parts`."""
    rows = table.rows(*key)
    changes = sorted({r.valid_from for r in rows} | {r.valid_to + timedelta(days=1) for r in rows if r.valid_to})
    result: list[tuple[date, date, TariffRow]] = []
    for i, start in enumerate(changes):
        end = changes[i + 1] - timedelta(days=1) if i + 1 < len(changes) else OPEN_END
        row = table.find_row(*key, start)
        if row is None:
            continue
        if result and result[-1][2] is row and result[-1][1] + timedelta(days=1) == start:
            result[-1] = (result[-1][0], end, row)
        else:
            result.append((start, end, row))
    return result


def _iso(expr: str) -> str:
    """Formula text for a date as YYYY-MM-DD, independent of the Excel language."""
    return f'YEAR({expr})&"-"&RIGHT("0"&MONTH({expr}),2)&"-"&RIGHT("0"&DAY({expr}),2)'


def _dates(first: str, last: str) -> str:
    """Formula text like the app's `dates_text`: one date, or "first to last"."""
    return f'{_iso(first)}&IF({first}={last},""," to "&{_iso(last)})'


class _Names:
    def __init__(self, wb: Workbook):
        self.wb = wb

    def add(self, name: str, ref: str) -> str:
        self.wb.defined_names[name] = DefinedName(name, attr_text=ref)
        return name


def _abs(sheet: str, col: int, row: int, col2: int | None = None, row2: int | None = None) -> str:
    ref = f"'{sheet}'!${get_column_letter(col)}${row}"
    if col2 is not None:
        ref += f":${get_column_letter(col2)}${row2}"
    return ref


# --- the workbook -----------------------------------------------------------------------------

def build_workbook(table: TariffTable, catalog: Catalog, source_name: str, today: date) -> Workbook:
    wb = Workbook()
    names = _Names(wb)
    calculator = wb.active
    calculator.title = "Calculator"
    calc = wb.create_sheet("Calc")
    lists = wb.create_sheet("Lists")
    points_ws = wb.create_sheet("Points")
    tariffs_ws = wb.create_sheet("Tariffs")
    rules = wb.create_sheet("Rules")

    _write_rules(rules, names)
    point_ids, slots, first_gy = _write_points_and_tariffs(table, catalog, points_ws, tariffs_ws, names)
    blocks = _write_lists(catalog, point_ids, lists, names)
    _write_calc(calc, names, blocks, slots, first_gy, today)
    _write_calculator(calculator, names, blocks, source_name, today)

    for ws in (calc, lists, points_ws, tariffs_ws, rules):
        ws.sheet_state = "hidden"
    wb.active = 0
    return wb


def _header(ws, headers: list[str]) -> None:
    for col, text in enumerate(headers, 1):
        ws.cell(1, col, text).font = HEADER_FONT


def _write_rules(ws, names: _Names) -> None:
    """One row per TSO in scope: the rules the formulas look up."""
    _header(ws, ["TSO code", "TSO", "Currency", "Unit", "Currency decimals", "Tariff decimals",
                 "Equal months: gas year", "Equal months: quarter", "Hours per unit"])
    for row, (code, (name, currency, unit)) in enumerate(TSOS.items(), 2):
        equal = fee.EQUAL_MONTH_INSTRUMENTS.get(code, ())
        values = [code, name, currency, unit, money.CURRENCY_DECIMALS[currency], fee.TARIFF_DECIMALS.get(code),
                  "yes" if fee.GAS_YEAR in equal else "no", "yes" if fee.QUARTER in equal else "no",
                  HOURS_PER_UNIT[unit]]
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
    last = len(TSOS) + 1
    for col, name in enumerate(["rl_code", "rl_name", "rl_currency", "rl_unit", "rl_curdec", "rl_tdec",
                                "rl_eq_gy", "rl_eq_q", "rl_hours"], 1):
        names.add(name, _abs(ws.title, col, 2, col, last))


def _write_points_and_tariffs(table: TariffTable, catalog: Catalog, points_ws, tariffs_ws, names: _Names):
    """The `Points` sheet (one row per catalog point, id = its position) and the
    `Tariffs` sheet (each point's stretches, contiguous and in date order)."""
    all_points: list[Point] = [*catalog.singles, *catalog.lng]
    _header(points_ws, ["Id", "TSO code", "TSO", "Direction", "EIC", "Name", "Capacity type", "Currency",
                        "Unit", "First Tariffs row", "Stretches", "Earliest gas year"])
    _header(tariffs_ws, ["Point id", "First day", "Last day", "Valid from", "Open-ended", "Valid to",
                         *periods.PRICE_COLUMNS])
    point_ids: dict[tuple[str, str, str], int] = {}
    tariff_row = 2
    most = 0
    first_gy = None
    for pid, point in enumerate(all_points, 1):
        parts = stretches(table, point.key)
        if not parts:
            raise WorkbookError(f"{point.tso_name} {point.cleaned_name} has no tariff rows.")
        most = max(most, len(parts))
        row = parts[0][2]
        earliest = periods.gas_year_start_year(table.earliest_valid_from(*point.key))
        first_gy = earliest if first_gy is None else min(first_gy, earliest)
        values = [pid, point.tso, point.tso_name, point.direction, point.eic, point.cleaned_name,
                  catalog.capacity_type(point), row.currency, row.unit, tariff_row, len(parts), earliest]
        for col, value in enumerate(values, 1):
            points_ws.cell(pid + 1, col, value)
        for first, last, part_row in parts:
            cells = [pid, first, last, part_row.valid_from, 1 if part_row.open_ended else 0, part_row.valid_to,
                     *(float(part_row.prices[c]) for c in periods.PRICE_COLUMNS)]
            for col, value in enumerate(cells, 1):
                cell = tariffs_ws.cell(tariff_row, col, value)
                if isinstance(value, date):
                    cell.number_format = "yyyy-mm-dd"
            tariff_row += 1
        point_ids[point.key] = pid
    if most > MAX_STRETCH_SLOTS:
        raise WorkbookError(f"A point has {most} tariff stretches; the workbook has room for {MAX_STRETCH_SLOTS}.")
    last = len(all_points) + 1
    for col, name in enumerate(["pt_id", "pt_tso", "pt_tso_name", "pt_direction", "pt_eic", "pt_name",
                                "pt_captype", "pt_currency", "pt_unit", "pt_frow", "pt_count", "pt_first_gy"], 1):
        names.add(name, _abs(points_ws.title, col, 2, col, last))
    names.add("tr_header", _abs(tariffs_ws.title, 7, 1, 6 + len(periods.PRICE_COLUMNS), 1))
    return point_ids, max(MIN_STRETCH_SLOTS, most), first_gy


def _write_lists(catalog: Catalog, point_ids: dict, ws, names: _Names) -> dict:
    """Dropdown lists. Column J holds every point label, each booking type (and
    each TSO's singles) as one block; K and L give the side ids, M the heading."""
    _header(ws, ["Booking type", "Instrument", "Unit", "", "TSO", "TSO code", "Singles start", "Singles count",
                 "", "Point label", "Side 1 id", "Side 2 id", "Heading"])
    for row, value in enumerate(BOOKING_TYPES, 2):
        ws.cell(row, 1, value)
    for row, instrument in enumerate(fee.INSTRUMENTS, 2):
        ws.cell(row, 2, fee.INSTRUMENT_LABELS[instrument])
    for row, unit in enumerate(fee.CAPACITY_UNITS, 2):
        ws.cell(row, 3, unit)
    names.add("lst_booking", _abs(ws.title, 1, 2, 1, 1 + len(BOOKING_TYPES)))
    names.add("lst_instrument", _abs(ws.title, 2, 2, 2, 1 + len(fee.INSTRUMENTS)))
    names.add("lst_unit", _abs(ws.title, 3, 2, 3, 1 + len(fee.CAPACITY_UNITS)))

    row = 2

    def block(entries) -> tuple[int, int]:
        nonlocal row
        start = row
        for label, id1, id2, heading in entries:
            for col, value in zip((10, 11, 12, 13), (label, id1, id2, heading)):
                ws.cell(row, col, value)
            row += 1
        return start, row - start

    blocks = {
        "Bundled": block((b.label(), point_ids[b.exit.key], point_ids[b.entry.key], b.label())
                         for b in catalog.bundles),
        "LNG entry": block((p.lng_label(), point_ids[p.key], None, p.lng_label()) for p in catalog.lng),
    }
    tsos = catalog.tsos()
    for i, tso in enumerate(tsos, 2):
        start, count = block((p.fee_label(catalog.capacity_type(p)), point_ids[p.key], None,
                              f"{tso_name(tso)} {p.fee_label()}") for p in catalog.singles_of(tso))
        for col, value in zip((5, 6, 7, 8), (tso_name(tso), tso, start, count)):
            ws.cell(i, col, value)
    last_tso = 1 + len(tsos)
    names.add("lst_tso_name", _abs(ws.title, 5, 2, 5, last_tso))
    names.add("lst_tso_start", _abs(ws.title, 7, 2, 7, last_tso))
    names.add("lst_tso_count", _abs(ws.title, 8, 2, 8, last_tso))
    names.add("lst_label", _abs(ws.title, 10, 1, 10, row - 1))
    names.add("lst_id1", _abs(ws.title, 11, 1, 11, row - 1))
    names.add("lst_id2", _abs(ws.title, 12, 1, 12, row - 1))
    names.add("lst_heading", _abs(ws.title, 13, 1, 13, row - 1))

    currencies = sorted({currency for _, currency, _ in TSOS.values()} - {money.EUR})
    for i, currency in enumerate(currencies, 2):
        ws.cell(i, 15, currency)
    ws.cell(1, 15, "Rate currency").font = HEADER_FONT
    names.add("lst_rate_cur", _abs(ws.title, 15, 2, 15, 1 + len(currencies)))
    blocks["rate_currencies"] = currencies
    return blocks


class _Calc:
    """Writes labelled scalars into columns A/B of the `Calc` sheet, each with a name."""

    def __init__(self, ws, names: _Names, row: int):
        self.ws, self.names, self.row = ws, names, row

    def scalar(self, name: str, formula: str) -> str:
        self.ws.cell(self.row, 1, name)
        self.ws.cell(self.row, 2, f"={formula}")
        self.names.add(name, _abs(self.ws.title, 2, self.row))
        self.row += 1
        return name


def _write_calc(ws, names: _Names, blocks: dict, slots: int, first_gy: int, today: date) -> None:
    ws.cell(1, 1, "Name").font = HEADER_FONT
    ws.cell(1, 2, "Value").font = HEADER_FONT
    ws.cell(1, 4, "Period list").font = HEADER_FONT
    c = _Calc(ws, names, 2)

    bundled_start, bundled_count = blocks["Bundled"]
    lng_start, lng_count = blocks["LNG entry"]
    c.scalar("c_booking_ok", "ISNUMBER(MATCH(in_booking,lst_booking,0))")
    c.scalar("c_bundle", 'in_booking="Bundled"')
    c.scalar("c_single", 'in_booking="Single side"')
    c.scalar("c_lng", 'in_booking="LNG entry"')
    c.scalar("c_tso_row", "IFERROR(MATCH(in_tso,lst_tso_name,0),0)")
    c.scalar("c_tso_ok", "OR(NOT(c_single),c_tso_row>0)")
    c.scalar("c_list_start", f"IF(c_bundle,{bundled_start},IF(c_lng,{lng_start},"
                             "IF(c_tso_row>0,INDEX(lst_tso_start,c_tso_row),1)))")
    c.scalar("c_list_count", f"IF(c_bundle,{bundled_count},IF(c_lng,{lng_count},"
                             "IF(c_tso_row>0,INDEX(lst_tso_count,c_tso_row),1)))")
    names.add("dv_points", "OFFSET(Lists!$J$1,c_list_start-1,0,MAX(1,c_list_count),1)")
    c.scalar("c_point_pos", "IF(AND(c_booking_ok,c_tso_ok),IFERROR(MATCH(in_point,dv_points,0),0),0)")
    c.scalar("c_point_ok", "c_point_pos>0")
    c.scalar("c_label_row", "c_list_start+c_point_pos-1")
    c.scalar("c_id1", 'IF(c_point_ok,INDEX(lst_id1,c_label_row),"")')
    c.scalar("c_id2", 'IF(AND(c_point_ok,c_bundle),INDEX(lst_id2,c_label_row),"")')
    c.scalar("c_heading_label", 'IF(c_point_ok,INDEX(lst_heading,c_label_row),"")')

    c.scalar("c_inst", "IFERROR(MATCH(in_instrument,lst_instrument,0),0)")
    c.scalar("c_today_gy", "IF(MONTH(TODAY())>=10,YEAR(TODAY()),YEAR(TODAY())-1)")
    c.scalar("c_last_gy", "c_today_gy+2")
    c.scalar("c_first_gy", "IF(c_point_ok,MIN(c_last_gy,MAX(INDEX(pt_first_gy,c_id1),"
                           "IF(c_bundle,INDEX(pt_first_gy,c_id2),0))),c_today_gy)")
    c.scalar("c_n_gy", "c_last_gy-c_first_gy+1")
    c.scalar("c_period_count", "IF(c_inst=1,c_n_gy,IF(c_inst=2,4*(c_n_gy+1),IF(c_inst=3,12*(c_n_gy+1),0)))")
    # Room for every gas year from the earliest tariff to 15 years after the build.
    period_slots = 12 * (periods.gas_year_start_year(today) + 15 - first_gy + 2)
    for k in range(1, period_slots + 1):
        gas_year = f'(c_first_gy+{k - 1})&"/"&RIGHT(c_first_gy+{k},2)'
        quarter = f'(c_first_gy+INT({k - 1}/4))&" Q"&(MOD({k - 1},4)+1)'
        month = f'(c_first_gy+INT({k - 1}/12))&"-"&RIGHT("0"&(MOD({k - 1},12)+1),2)'
        ws.cell(1 + k, 4, f'=IF({k}>c_period_count,"",IF(c_inst=1,{gas_year},IF(c_inst=2,{quarter},{month})))')
    names.add("dv_periods", f"OFFSET(Calc!$D$2,0,0,MAX(1,MIN({period_slots},c_period_count)),1)")
    c.scalar("c_period_pos", 'IF(AND(c_inst>=1,c_inst<=3,c_period_count>0),'
                             'IFERROR(MATCH(in_period&"",dv_periods,0),0),0)')
    c.scalar("c_day_ok", "AND(c_inst=4,ISNUMBER(in_day))")
    c.scalar("c_period_ok", "IF(c_inst=4,c_day_ok,c_period_pos>0)")
    c.scalar("c_py", "c_first_gy+IF(c_inst=2,INT((c_period_pos-1)/4),IF(c_inst=3,INT((c_period_pos-1)/12),c_period_pos-1))")
    c.scalar("c_pq", "MOD(c_period_pos-1,4)+1")
    c.scalar("c_pm", "MOD(c_period_pos-1,12)+1")
    c.scalar("c_start", 'IF(NOT(c_period_ok),"",CHOOSE(c_inst,DATE(c_py,10,1),DATE(c_py,3*c_pq-2,1),'
                        'DATE(c_py,c_pm,1),INT(in_day)))')
    c.scalar("c_end", 'IF(NOT(c_period_ok),"",CHOOSE(c_inst,DATE(c_py+1,9,30),DATE(c_py,3*c_pq+1,0),'
                      'DATE(c_py,c_pm+1,0),INT(in_day)))')
    c.scalar("c_unit_hours", 'IF(in_unit="kWh/h",24,IF(in_unit="kWh/d",1,0))')
    c.scalar("c_capacity_ok", "AND(ISNUMBER(in_capacity),IF(ISNUMBER(in_capacity),in_capacity>0,FALSE))")
    c.scalar("c_kwh_d", "IF(c_capacity_ok,in_capacity*c_unit_hours,0)")
    c.scalar("c_input_msg",
             'IF(NOT(c_booking_ok),"Choose a booking type from the list.",'
             'IF(NOT(c_tso_ok),"Choose a TSO from the list.",'
             'IF(NOT(c_point_ok),"Choose a point from the list.",'
             'IF(c_inst=0,"Choose a product instrument from the list.",'
             'IF(NOT(c_period_ok),IF(c_inst=4,"Choose a product period.","Choose a period from the list."),'
             'IF(c_unit_hours=0,"Choose a unit from the list.",'
             'IF(NOT(c_capacity_ok),"' + fee.CAPACITY_MESSAGE + '","")))))))')
    c.scalar("c_inputs_ok", 'c_input_msg=""')
    c.scalar("c_notariff_msg", 'IF(sa_notariff,sa_missing_msg,IF(sb_notariff,sb_missing_msg,""))')
    c.scalar("c_msg", 'IF(c_input_msg<>"",c_input_msg,c_notariff_msg)')
    c.scalar("c_ok", 'c_msg=""')
    c.scalar("c_heading", f'IF(c_ok,c_heading_label&": "&LOWER(in_instrument)&" "&{_dates("c_start", "c_end")},"")')

    # The bundle's rates and EUR total.
    c.scalar("c_eur_on", "AND(c_ok,c_bundle)")
    c.scalar("c_r1_cur", 'IF(AND(c_eur_on,sa_currency<>"EUR"),sa_currency,"")')
    c.scalar("c_r2_cur", 'IF(AND(c_eur_on,sb_currency<>"EUR",sb_currency<>sa_currency),sb_currency,"")')
    for i in (1, 2):
        cur = f"c_r{i}_cur"
        cell = f"INDEX(in_rates,MATCH({cur},lst_rate_cur,0))"
        c.scalar(f"c_r{i}_state", f'IF({cur}="","",IF(ISBLANK({cell}),"missing",'
                                  f'IF(ISNUMBER({cell}),IF({cell}>0,"ok","invalid"),"invalid")))')
        c.scalar(f"c_r{i}_rate", f'IF(c_r{i}_state="ok",{cell},"")')
    for state, name in (("missing", "c_missing"), ("invalid", "c_invalid")):
        one = f'IF(c_r1_state="{state}",c_r1_cur&" per EUR","")'
        two = f'IF(c_r2_state="{state}",c_r2_cur&" per EUR","")'
        c.scalar(f"{name}_1", one)
        c.scalar(f"{name}_2", two)
        c.scalar(f"{name}_names", f'IF(AND({name}_1<>"",{name}_2<>""),{name}_1&" and "&{name}_2,{name}_1&{name}_2)')
    c.scalar("c_missing_text", 'IF(c_missing_names="","","Enter the "&c_missing_names&" rate"'
                               '&IF(AND(c_missing_1<>"",c_missing_2<>""),"s","")&" to see the EUR total.")')
    c.scalar("c_invalid_text", 'IF(c_invalid_names="","","The "&c_invalid_names&" rate must be a positive number.")')
    c.scalar("c_eur_msg", 'IF(NOT(c_eur_on),"",IF(AND(c_missing_text<>"",c_invalid_text<>""),'
                          'c_missing_text&" "&c_invalid_text,c_missing_text&c_invalid_text))')
    c.scalar("c_eur_ready", 'AND(c_eur_on,c_eur_msg="")')
    side_rate = lambda n: f'IF({SIDE_PREFIX[n]}currency="EUR",1,IF({SIDE_PREFIX[n]}currency=c_r1_cur,c_r1_rate,c_r2_rate))'
    c.scalar("c_eur_total", f'IF(c_eur_ready,ROUND(ROUND(sa_total/{side_rate(1)}+sb_total/{side_rate(2)},8),2),"")')
    for i in (1, 2):
        c.scalar(f"c_r{i}_text", f'IF(c_r{i}_cur="","","at 1 EUR = "&c_r{i}_rate&" "&c_r{i}_cur)')
    c.scalar("c_eur_caption", 'IF(NOT(c_eur_ready),"",IF(AND(c_r1_cur="",c_r2_cur=""),'
                              '"Sum of both sides\' totals, both in EUR.","Sum of both sides\' totals, converted "'
                              '&c_r1_text&IF(AND(c_r1_text<>"",c_r2_text<>""),"; ","")&c_r2_text&"."))')
    for currency in blocks["rate_currencies"]:
        c.scalar(f"c_need_{currency}", f'AND(c_eur_on,OR(sa_currency="{currency}",sb_currency="{currency}"))')

    base = max(c.row, period_slots + 2) + 5
    for n in (1, 2):
        _write_side(ws, names, n, base + (n - 1) * (max(slots, 12, 45) + 10), slots)


def _write_side(ws, names: _Names, n: int, base: int, slots: int) -> None:
    """One side's pricing: scalars in A/B, the stretch slots, the month slots and
    the overlap grid between them."""
    s = SIDE_PREFIX[n]
    c = _Calc(ws, names, base)
    k_rows = range(base + 1, base + 1 + slots)  # stretch slots
    m_rows = range(base + 1, base + 13)  # month slots
    col = {name: i for i, name in enumerate(
        ["k", "present", "trow", "first", "last", "vfrom", "open", "raw", "cell", "overlap", "distinct",
         "ptext", "parttext", "miss"], 4)}
    mcol = {name: i for i, name in enumerate(
        ["j", "present", "first", "last", "days", "price", "wraw", "w", "share", "floor", "frac", "rank",
         "minor", "label", "amount"], 19)}
    grid0 = 35  # first column of the stretch x month overlap grid
    L = get_column_letter

    def k_rng(name: str) -> str:
        return f"${L(col[name])}${k_rows[0]}:${L(col[name])}${k_rows[-1]}"

    def m_rng(name: str) -> str:
        return f"${L(mcol[name])}${m_rows[0]}:${L(mcol[name])}${m_rows[-1]}"

    c.scalar(s + "id", "c_id1" if n == 1 else "c_id2")
    c.scalar(s + "has", f"ISNUMBER({s}id)")
    c.scalar(s + "ready", f"AND(c_inputs_ok,{s}has)")
    for field, rng in (("tso", "pt_tso"), ("tso_name", "pt_tso_name"), ("direction", "pt_direction"),
                       ("name", "pt_name"), ("currency", "pt_currency"), ("unit", "pt_unit"),
                       ("captype", "pt_captype"), ("frow", "pt_frow"), ("count", "pt_count")):
        c.scalar(s + field, f'IF({s}has,INDEX({rng},{s}id),"")')
    c.scalar(s + "rule", f"IFERROR(MATCH({s}tso,rl_code,0),0)")
    c.scalar(s + "curdec", f"IF({s}rule>0,INDEX(rl_curdec,{s}rule),2)")
    c.scalar(s + "tdec", f'IF({s}rule=0,"",IF(INDEX(rl_tdec,{s}rule)="","",INDEX(rl_tdec,{s}rule)))')
    c.scalar(s + "equal", f'IF({s}rule=0,FALSE,IF(c_inst=1,INDEX(rl_eq_gy,{s}rule)="yes",'
                          f'IF(c_inst=2,INDEX(rl_eq_q,{s}rule)="yes",FALSE)))')
    c.scalar(s + "uhours", f'IF({s}unit="kWh/d",1,24)')
    c.scalar(s + "col", f"IF(NOT({s}ready),1,CHOOSE(c_inst,1,2+INT((MONTH(c_start)-1)/3),"
                        f"5+MONTH(c_start),17+MONTH(c_start)))")
    c.scalar(s + "colname", f"INDEX(tr_header,{s}col)")
    c.scalar(s + "pdays", f"IF({s}ready,c_end-c_start+1,0)")
    c.scalar(s + "nmonths", f"IF({s}ready,(YEAR(c_end)-YEAR(c_start))*12+MONTH(c_end)-MONTH(c_start)+1,0)")
    c.scalar(s + "covered", f"SUM({k_rng('overlap')})")
    c.scalar(s + "notariff", f"AND({s}ready,{s}covered<{s}pdays)")
    c.scalar(s + "miss_start", f'IF({s}notariff,IF(SUMPRODUCT(({k_rng("first")}<=c_start)*({k_rng("last")}>=c_start)'
                               f'*{k_rng("present")})=0,c_start,""),"")')
    c.scalar(s + "firstmiss", f'IF({s}notariff,MIN({s}miss_start,{k_rng("miss")}),"")')
    if n == 1:
        side_text = f'{s}tso_name&" "&{s}name'
        bundle_text = f'"the "&{s}tso_name&" "&IF({s}direction="{EXIT}","exit","entry")&" side ("&{s}name&")"'
        c.scalar(s + "missing_msg", f'IF({s}notariff,"No tariff exists for "&IF(c_bundle,{bundle_text},{side_text})'
                                    f'&" on "&{_iso(s + "firstmiss")}&".","")')
    else:
        bundle_text = f'"the "&{s}tso_name&" "&IF({s}direction="{EXIT}","exit","entry")&" side ("&{s}name&")"'
        c.scalar(s + "missing_msg", f'IF({s}notariff,"No tariff exists for "&{bundle_text}'
                                    f'&" on "&{_iso(s + "firstmiss")}&".","")')
    c.scalar(s + "show", f"AND(c_ok,{s}has)")
    c.scalar(s + "nparts", f'COUNTIF({k_rng("overlap")},">0")')
    c.scalar(s + "price", f'IF(AND({s}ready,NOT({s}notariff)),IF({s}equal,SUM({m_rng("price")})/{s}nmonths,'
                          f'SUMPRODUCT({m_rng("price")},{m_rng("days")})/{s}pdays),0)')
    c.scalar(s + "capacity", f"c_kwh_d/{s}uhours")
    c.scalar(s + "total", f"ROUND(ROUND(c_kwh_d*{s}price/{s}uhours,{s}curdec+6),{s}curdec)")
    c.scalar(s + "total_minor", f"ROUND({s}total*10^{s}curdec,0)")
    c.scalar(s + "wraw_sum", f"SUM({m_rng('wraw')})")
    c.scalar(s + "wsum", f"SUM({m_rng('w')})")
    c.scalar(s + "leftover", f"{s}total_minor-SUM({m_rng('floor')})")
    c.scalar(s + "price_text", f'_xlfn.TEXTJOIN(" / ",TRUE,{k_rng("ptext")})')
    c.scalar(s + "tariff_used", f'_xlfn.TEXTJOIN("; ",TRUE,{k_rng("parttext")})')
    c.scalar(s + "heading", f'IF(c_bundle,{s}tso_name&" "&LOWER({s}direction)&" "&{s}name,"")')

    for name, i in col.items():
        ws.cell(base, i, f"{s}{name}").font = HEADER_FONT
    for name, i in mcol.items():
        ws.cell(base, i, f"{s}m_{name}").font = HEADER_FONT

    def ref(name: str, row: int, table=col) -> str:
        return f"{L(table[name])}{row}"

    first_k = k_rows[0]
    for k, row in enumerate(k_rows, 1):
        r = lambda name: ref(name, row)
        above = lambda name: f"${L(col[name])}${first_k}:${L(col[name])}{row - 1}"
        cells = {
            "k": k,
            "present": f"=IF(AND({s}has,{k}<={s}count),1,0)",
            "trow": f"=IF({r('present')}=1,{s}frow+{k - 1},0)",
            "first": f"=IF({r('present')}=1,INDEX(Tariffs!$B:$B,{r('trow')}),0)",
            "last": f"=IF({r('present')}=1,INDEX(Tariffs!$C:$C,{r('trow')}),0)",
            "vfrom": f"=IF({r('present')}=1,INDEX(Tariffs!$D:$D,{r('trow')}),0)",
            "open": f"=IF({r('present')}=1,INDEX(Tariffs!$E:$E,{r('trow')})=1,FALSE)",
            "raw": f"=IF({r('present')}=1,INDEX(Tariffs!$G:$AI,{r('trow')},{s}col),0)",
            "cell": f'=IF({s}tdec="",{r("raw")},ROUND({r("raw")},{s}tdec))',
            "overlap": f"=IF(AND({s}ready,{r('present')}=1),MAX(0,MIN({r('last')},c_end)-MAX({r('first')},c_start)+1),0)",
            "distinct": (f"={r('overlap')}>0" if k == 1 else
                         f"=AND({r('overlap')}>0,SUMPRODUCT(({above('cell')}={r('cell')})*({above('overlap')}>0))=0)"),
            "ptext": f'=IF({r("distinct")},FIXED({r("cell")},4,TRUE),"")',
            "parttext": (f'=IF({r("overlap")}>0,"from "&{_iso(r("vfrom"))}&IF({r("open")}," (open-ended)","")'
                         f'&IF({s}nparts>1," for "&{_dates(f"MAX({r('first')},c_start)", f"MIN({r("last")},c_end)")},""),"")'),
            "miss": (f'=IF({s}ready,IF(AND({r("present")}=1,{r("last")}+1>=c_start,{r("last")}+1<=c_end,'
                     f'SUMPRODUCT(({k_rng("first")}<={r("last")}+1)*({k_rng("last")}>={r("last")}+1)*{k_rng("present")})=0),'
                     f'{r("last")}+1,""),"")'),
        }
        for name, value in cells.items():
            ws.cell(row, col[name], value)

    first_m = m_rows[0]
    for j, row in enumerate(m_rows, 1):
        r = lambda name: ref(name, row, mcol)
        grid_col = L(grid0 + j - 1)
        cells = {
            "j": j,
            "present": f"=IF(AND({s}ready,{j}<={s}nmonths),1,0)",
            "first": f"=IF({r('present')}=1,MAX(c_start,DATE(YEAR(c_start),MONTH(c_start)+{j - 1},1)),0)",
            "last": f"=IF({r('present')}=1,MIN(c_end,DATE(YEAR(c_start),MONTH(c_start)+{j},0)),0)",
            "days": f"=IF({r('present')}=1,{r('last')}-{r('first')}+1,0)",
            "price": (f"=IF({r('present')}=1,SUMPRODUCT({k_rng('cell')},"
                      f"${grid_col}${k_rows[0]}:${grid_col}${k_rows[-1]})/{r('days')},0)"),
            "wraw": f"=IF({r('present')}=1,IF({s}equal,{r('price')},{r('price')}*{r('days')}),0)",
            "w": f"=IF({s}wraw_sum=0,{r('days')},{r('wraw')})",
            "share": f"=IF(AND({r('present')}=1,{s}wsum>0),ROUND({s}total_minor*{r('w')}/{s}wsum,6),0)",
            "floor": f"=INT({r('share')})",
            "frac": f"=IF({r('present')}=1,ROUND({r('share')}-{r('floor')},6),-1)",
            "rank": (f"=IF({r('present')}=1,1+SUMPRODUCT(({m_rng('frac')}>{r('frac')})*1)"
                     + ("" if j == 1 else
                        f"+SUMPRODUCT((${L(mcol['frac'])}${first_m}:${L(mcol['frac'])}{row - 1}={r('frac')})*1)")
                     + ",0)"),
            "minor": f"=IF({r('present')}=1,{r('floor')}+IF({r('rank')}<={s}leftover,1,0),0)",
            "label": f'=IF({r("present")}=1,YEAR({r("first")})&"-"&RIGHT("0"&MONTH({r("first")}),2),"")',
            "amount": f'=IF({r("present")}=1,{r("minor")}/10^{s}curdec,"")',
        }
        for name, value in cells.items():
            ws.cell(row, mcol[name], value)
        ws.cell(base, grid0 + j - 1, f"{s}grid_{j}").font = HEADER_FONT
        for row_k in k_rows:
            ws.cell(row_k, grid0 + j - 1,
                    f"=IF(AND({ref('present', row_k)}=1,${L(mcol['present'])}${row}=1),"
                    f"MAX(0,MIN({ref('last', row_k)},${L(mcol['last'])}${row})"
                    f"-MAX({ref('first', row_k)},${L(mcol['first'])}${row})+1),0)")

    for name in ("label", "days", "amount"):
        names.add(f"{s}m_{name}", f"Calc!{m_rng(name)}")
    names.add(f"{s}m_present", f"Calc!{m_rng('present')}")




def _write_calculator(ws, names: _Names, blocks: dict, source_name: str, today: date) -> None:
    """The one visible sheet: inputs on top, then each side's results, then the bundle's EUR total."""
    for column, width in zip("ABCD", (30, 60, 16, 44)):
        ws.column_dimensions[column].width = width
    ws["A1"] = "Capacity Fee Calculator"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"Built from {source_name} on {today.isoformat()}. Fill in the yellow cells."
    ws["A2"].font = NOTE_FONT

    def put(cell: str, value, name: str | None = None, number_format: str | None = None, font=None) -> None:
        ws[cell] = value
        if name:
            column = "".join(ch for ch in cell if ch.isalpha())
            names.add(name, f"Calculator!${column}${cell[len(column):]}")
        if number_format:
            ws[cell].number_format = number_format
        if font:
            ws[cell].font = font

    def grey_when(cell: str, condition: str) -> None:
        ws.conditional_formatting.add(cell, Rule(type="expression", dxf=GREY, formula=[condition]))

    def list_input(cell: str, name: str, source: str, default) -> None:
        put(cell, default, name)
        dv = DataValidation(type="list", formula1=source, allow_blank=True, showErrorMessage=True,
                            errorTitle="Not in the list", error="Choose a value from the dropdown list.")
        ws.add_data_validation(dv)
        dv.add(cell)

    ws["A4"] = "Booking type"
    list_input("B4", "in_booking", "lst_booking", BOOKING_TYPES[0])
    ws["A5"] = "TSO"
    list_input("B5", "in_tso", "lst_tso_name", tso_name(DEFAULT_TSO))
    put("D5", '=IF(c_single,"","Only used for Single side")', font=NOTE_FONT)
    grey_when("B5", "NOT(c_single)")
    ws["A6"] = '=IF(c_bundle,"Bundled point",IF(c_lng,"LNG terminal","Network point"))'
    list_input("B6", "in_point", "dv_points", DEFAULT_BUNDLE)
    ws["A7"] = "Product instrument"
    list_input("B7", "in_instrument", "lst_instrument", fee.INSTRUMENT_LABELS[fee.QUARTER])
    ws["A8"] = "Period"
    list_input("B8", "in_period", "dv_periods", DEFAULT_PERIOD)
    ws["B8"].number_format = "@"  # keep "2026-10" as text, never a date
    put("D8", '=IF(c_inst=4,"Not used for a day: enter the day below","")', font=NOTE_FONT)
    grey_when("B8", "c_inst=4")
    ws["A9"] = "Day"
    put("B9", DEFAULT_DAY, "in_day", "yyyy-mm-dd")
    day_dv = DataValidation(type="date", operator="between", formula1="36526", formula2="73415",
                            allow_blank=True, showErrorMessage=True, errorTitle="Not a date",
                            error="Enter a date between 2000-01-01 and 2100-12-31.")
    ws.add_data_validation(day_dv)
    day_dv.add("B9")
    put("D9", '=IF(c_inst=4,"","Only used for the Day instrument")', font=NOTE_FONT)
    grey_when("B9", "c_inst<>4")
    ws["A10"] = "Requested capacity"
    put("B10", None, "in_capacity")
    list_input("C10", "in_unit", "lst_unit", fee.CAPACITY_UNITS[0])
    inputs = ["B4", "B5", "B6", "B7", "B8", "B9", "B10", "C10"]
    for cell in ("A4", "A5", "A6", "A7", "A8", "A9", "A10"):
        ws[cell].font = HEADER_FONT

    put("A12", "=c_msg", "msg_main", font=Font(bold=True, color="C55A11"))
    put("A13", "=c_heading", "out_heading", font=Font(bold=True, size=13))

    money_cells: list[tuple[str, str]] = []  # (cell, side prefix): HUF amounts show whole forints
    for n in (1, 2):
        s = SIDE_PREFIX[n]
        o = FIRST_SIDE_ROW + (n - 1) * SIDE_HEIGHT

        def when(text: str, extra: str = "") -> str:
            condition = f"AND({s}show,{extra})" if extra else f"{s}show"
            return f'=IF({condition},{text},"")'

        put(f"A{o}", when(f"{s}heading"), f"out{n}_heading", font=Font(bold=True, size=12))
        put(f"A{o + 1}", when('"TSO"'), font=HEADER_FONT)
        put(f"B{o + 1}", when(f"{s}tso_name"), f"out{n}_tso")
        put(f"C{o + 1}", when('"Direction"'), font=HEADER_FONT)
        put(f"D{o + 1}", when(f"{s}direction"), f"out{n}_direction")
        put(f"A{o + 2}", when(f'"Capacity ("&{s}unit&")"'), font=HEADER_FONT)
        put(f"B{o + 2}", when(f"{s}capacity"), f"out{n}_capacity", "#,##0.00")
        put(f"A{o + 3}", when(f'"Price ("&{s}colname&", "&{s}currency&" per "&{s}unit&")"'), font=HEADER_FONT)
        put(f"B{o + 3}", when(f"{s}price_text"), f"out{n}_price")
        put(f"A{o + 4}", when(f'"Total cost ("&{s}currency&")"'), font=HEADER_FONT)
        put(f"B{o + 4}", when(f"{s}total"), f"out{n}_total", "#,##0.00", Font(bold=True))
        money_cells.append((f"B{o + 4}", s))
        put(f"A{o + 5}", when('"Tariff used"'), font=HEADER_FONT)
        put(f"B{o + 5}", when(f"{s}tariff_used"), f"out{n}_tariff_used")
        put(f"A{o + 6}", when(f'"All amounts are in "&{s}currency&"; the total is rounded to "'
                              f'&IF({s}curdec=0,"whole forints","cents")&"."'), f"out{n}_caption", font=NOTE_FONT)
        put(f"A{o + 7}", when(f'{s}tso_name&" bills its tariff rounded to "&{s}tdec&" decimals."', f'{s}tdec<>""'),
            f"out{n}_rounding", font=NOTE_FONT)
        put(f"A{o + 8}", when('"This point has no Firm tariff, so Interruptible capacity is priced; '
                              'the TSO can interrupt it."', f'{s}captype="{INTERRUPTIBLE}"'),
            f"out{n}_interruptible", font=Font(bold=True, color="C00000"))
        title = 'IF(c_bundle,"Monthly invoice ("&sa_tso_name&")","Monthly invoice")' if n == 1 else \
            '"Monthly invoice ("&sb_tso_name&")"'
        put(f"A{o + 9}", when(title), font=HEADER_FONT)
        put(f"A{o + 10}", when(f'{s}tso_name&" invoices each month as "&IF(c_inst=1,"1/12 of the gas year",'
                               f'"1/3 of the quarter")&" at the tariff valid in it, whatever the month\'s number '
                               f'of days."', f"{s}equal"), f"out{n}_equal", font=NOTE_FONT)
        put(f"A{o + 11}", when('"Month"'), font=HEADER_FONT)
        put(f"B{o + 11}", when('"Days"'), font=HEADER_FONT)
        put(f"C{o + 11}", when(f'"Amount ("&{s}currency&")"'), font=HEADER_FONT)
        for j in range(1, 13):
            row = o + 11 + j
            present = f"INDEX({s}m_present,{j})=1"
            put(f"A{row}", when(f"INDEX({s}m_label,{j})", present))
            put(f"B{row}", when(f"INDEX({s}m_days,{j})", present))
            put(f"C{row}", when(f"INDEX({s}m_amount,{j})", present), number_format="#,##0.00")
            money_cells.append((f"C{row}", s))
        names.add(f"out{n}_months", f"Calculator!$A${o + 12}:$A${o + 23}")
        names.add(f"out{n}_days", f"Calculator!$B${o + 12}:$B${o + 23}")
        names.add(f"out{n}_amounts", f"Calculator!$C${o + 12}:$C${o + 23}")
        put(f"A{o + 24}", when('"Total"'), font=HEADER_FONT)
        put(f"B{o + 24}", when(f"{s}pdays"))
        put(f"C{o + 24}", when(f"{s}total"), f"out{n}_invoice_total", "#,##0.00", HEADER_FONT)
        money_cells.append((f"C{o + 24}", s))

    whole = DifferentialStyle(numFmt=NumberFormat(numFmtId=3, formatCode="#,##0"))
    for cell, s in money_cells:
        ws.conditional_formatting.add(cell, Rule(type="expression", dxf=whole, formula=[f'{s}currency="HUF"']))

    o = FIRST_SIDE_ROW + 2 * SIDE_HEIGHT
    put(f"A{o}", '=IF(c_bundle,"Bundle total in EUR","")', font=Font(bold=True, size=12))
    put(f"A{o + 1}", '=IF(c_bundle,"Exchange rates: only the ones this bundle needs are used.","")', font=NOTE_FONT)
    currencies = blocks["rate_currencies"]
    for i, currency in enumerate(currencies):
        row = o + 2 + i
        put(f"A{row}", f"{currency} per 1 EUR", font=HEADER_FONT)
        put(f"B{row}", None, f"in_rate_{currency}")
        put(f"D{row}", f'=IF(c_need_{currency},"","Not used for this booking")', font=NOTE_FONT)
        grey_when(f"B{row}", f"NOT(c_need_{currency})")
        inputs.append(f"B{row}")
    names.add("in_rates", f"Calculator!$B${o + 2}:$B${o + 1 + len(currencies)}")
    row = o + 2 + len(currencies)
    put(f"A{row}", '=IF(c_eur_ready,"Bundle total (EUR)","")', font=HEADER_FONT)
    put(f"B{row}", "=c_eur_total", "out_eur_total", "#,##0.00", Font(bold=True))
    put(f"A{row + 1}", "=c_eur_caption", "out_eur_caption", font=NOTE_FONT)
    put(f"A{row + 2}", "=c_eur_msg", "out_eur_msg", font=Font(bold=True, color="C55A11"))

    for cell in inputs:
        ws[cell].fill = INPUT_FILL
        ws[cell].protection = Protection(locked=False)
    ws.protection.sheet = True
    ws.print_area = f"A1:D{row + 2}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
